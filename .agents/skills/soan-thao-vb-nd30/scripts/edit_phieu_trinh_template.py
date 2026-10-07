#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
edit_phieu_trinh_template.py — Chuyên biệt hóa điền nội dung Phiếu trình mẫu
NGUYÊN TẮC BẢO TOÀN NGUYÊN BẢN 100% (ZERO STRUCTURE MUTATION):
- Giữ nguyên 100% khung viền, bảng lồng, cột, lề, font chữ Times New Roman của template gốc:
  bosung/Phiếu trình.docx
- Sử dụng Run-Level In-Place Replacement: tuyệt đối không xóa paragraph, không tạo mới docx từ đầu.
"""

import os
import sys
import docx

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

# Import core engine
from scripts.edit_docx_template_in_place import (
    replace_text_in_paragraph_exact,
    replace_after_prefix_in_paragraph
)

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "..", "bosung", "Phiếu trình.docx")


def fill_phieu_trinh(data, output_path, template_path=None):
    """
    Điền dữ liệu vào Phiếu trình bảo toàn 100% nguyên bản cấu trúc và font chữ.
    """
    if template_path is None:
        template_path = TEMPLATE_PATH

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Không tìm thấy template mẫu: {template_path}")

    doc = docx.Document(template_path)

    # 1. CẬP NHẬT TABLE 0 (HEADER - SỐ HIỆU)
    t0 = doc.tables[0]
    so_ky_hieu = data.get("so_ky_hieu")
    if so_ky_hieu:
        for p in t0.rows[0].cells[0].paragraphs:
            if "Số:" in p.text:
                replace_after_prefix_in_paragraph(p, "Số:", so_ky_hieu)

    # 2. CẬP NHẬT TABLE 2 (NỘI DUNG CHÍNH TRONG KHUNG)
    t2 = doc.tables[2]
    cell_main = t2.rows[0].cells[0]
    paragraphs = cell_main.paragraphs

    # Mục 1: Văn bản/cơ quan, đơn vị
    if data.get("van_ban_trinh"):
        for p in paragraphs:
            if "1. Văn bản/cơ quan, đơn vị:" in p.text:
                replace_after_prefix_in_paragraph(p, "1. Văn bản/cơ quan, đơn vị:", data["van_ban_trinh"])
                break

    # Mục 2: Trích yếu
    if data.get("trich_yeu"):
        for p in paragraphs:
            if "2. Trích yếu:" in p.text:
                replace_after_prefix_in_paragraph(p, "2. Trích yếu:", data["trich_yeu"])
                break

    # Mục 2 (dưới phần II): Các sở, ban, ngành trình
    if data.get("cac_co_quan_trinh"):
        for p in paragraphs:
            if "2. Các sở, ban, ngành" in p.text or "2. Các sở, ban, ngành" in p.text:
                prefix = "2. Các sở, ban, ngành, đơn vị, địa phương trình (kết quả đề nghị):"
                if prefix in p.text:
                    replace_after_prefix_in_paragraph(p, prefix, data["cac_co_quan_trinh"])
                else:
                    # Match linh hoạt
                    prefix_short = "2. Các sở, ban, ngành"
                    replace_after_prefix_in_paragraph(p, prefix_short, data["cac_co_quan_trinh"])
                break

    # Mục 3: Thẩm quyền
    if data.get("tham_quyen"):
        for p in paragraphs:
            if "3. Thẩm quyền" in p.text:
                prefix = "3. Thẩm quyền (Chọn ghi thẩm quyền tương ứng):"
                if prefix in p.text:
                    replace_after_prefix_in_paragraph(p, prefix, data["tham_quyen"])
                else:
                    replace_after_prefix_in_paragraph(p, "3. Thẩm quyền", data["tham_quyen"])
                break

    # Mục 4: Đề xuất phương án (paragraphs 12..16)
    de_xuat_text = data.get("de_xuat_phuong_an")
    if de_xuat_text:
        # P12 là đoạn đầu tiên của nội dung đề xuất
        # Đưa toàn bộ nội dung đề xuất mới vào P12, xóa sạch chữ ở P13..P16 nhưng giữ nguyên paragraph
        paragraphs[12].runs[0].text = de_xuat_text
        for r in paragraphs[12].runs[1:]:
            r.text = ""
        for pi in range(13, 17):
            for r in paragraphs[pi].runs:
                r.text = ""

    # Mục 5: Độ mật
    if data.get("do_mat"):
        for p in paragraphs:
            if "5. Độ mật:" in p.text:
                replace_after_prefix_in_paragraph(p, "5. Độ mật:", data["do_mat"])
                break

    # 3. CẬP NHẬT KHỐI CHỮ KÝ VÀ NGÀY THÁNG (TABLE 2 ROW 1, 2, 3)
    ngay_thang = data.get("ngay_thang", "Ngày 25  tháng 9 năm 2026")
    chuyen_vien = data.get("chuyen_vien", "Trương Hải Châu")
    lanh_dao_phong = data.get("lanh_dao_phong", "Lương Thị Nguyệt Thanh")
    lanh_dao_vp = data.get("lanh_dao_vp", "Trần Thanh Toàn")
    lanh_dao_tinh = data.get("lanh_dao_tinh", "Đinh Văn Tuấn")

    # Thay ngày tháng mẫu cũ: 'Ngày 03  tháng 12 năm 2025' hoặc các dạng tương tự
    exact_replacements = {
        "Ngày 03  tháng 12 năm 2025": ngay_thang,
        "Ngày  03 tháng 12 năm 2025": ngay_thang,
        "Nguyễn Tuấn Anh": lanh_dao_phong,
        "Phan Sỹ Thống": lanh_dao_vp,
        "Đinh Văn Tuấn": lanh_dao_tinh,
    }
    if chuyen_vien != "Trương Hải Châu":
        exact_replacements["Trương Hải Châu"] = chuyen_vien

    # Duyệt thay thế chính xác trong toàn bộ tài liệu (đặc biệt các ô chữ ký)
    for s_text, r_text in exact_replacements.items():
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if s_text in p.text:
                            replace_text_in_paragraph_exact(p, s_text, r_text)

    # Lưu file
    doc.save(output_path)
    print(f"✅ Đã điền và xuất Phiếu trình bảo toàn 100% cấu trúc dòng và font chữ: {output_path}")
    return output_path


if __name__ == "__main__":
    test_data = {
        "so_ky_hieu": "99-9/26/PT-NC",
        "van_ban_trinh": "Công văn số 4329/VPCP-CĐS ngày 14/9/2026 của Văn phòng Chính phủ.",
        "trich_yeu": "V/v báo cáo tình hình an ninh mạng và chuyển đổi số quý III/2026.",
        "cac_co_quan_trinh": "Sở Thông tin và Truyền thông phối hợp Công an tỉnh.",
        "tham_quyen": "Chủ tịch Ủy ban nhân dân tỉnh.",
        "de_xuat_phuong_an": "Chuyên viên kính đề xuất Chủ tịch UBND tỉnh giao Công an tỉnh chủ trì, phối hợp với Sở Khoa học và Công nghệ và các cơ quan liên quan triển khai thực hiện Công văn số 4329/VPCP-CĐS nêu trên theo đúng thẩm quyền.",
        "chuyen_vien": "Trương Hải Châu",
        "lanh_dao_phong": "Lương Thị Nguyệt Thanh",
        "lanh_dao_vp": "Trần Thanh Toàn",
        "lanh_dao_tinh": "Đinh Văn Tuấn",
        "ngay_thang": "Ngày 25  tháng 9 năm 2026",
        "do_mat": "Mật",
    }

    out_file = os.path.join(os.path.dirname(__file__), "test_phieu_trinh_exact_output.docx")
    fill_phieu_trinh(test_data, out_file)
