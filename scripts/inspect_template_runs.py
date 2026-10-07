import sys
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'e:\OneDrive\HSCV\Antigravity\vbaibot\bosung\Phiếu trình.docx')
t2 = doc.tables[2]
p1 = t2.rows[0].cells[0].paragraphs[1]
print('P1 full text:', p1.text)
print('P1 run count:', len(p1.runs))
for i, r in enumerate(p1.runs):
    print(f'  Run {i}: text="{r.text}", font={r.font.name}, size={r.font.size}, bold={r.bold}, italic={r.italic}')

p2 = t2.rows[0].cells[0].paragraphs[2]
print('\nP2 full text:', p2.text)
print('P2 run count:', len(p2.runs))
for i, r in enumerate(p2.runs):
    print(f'  Run {i}: text="{r.text}", font={r.font.name}, size={r.font.size}, bold={r.bold}, italic={r.italic}')
