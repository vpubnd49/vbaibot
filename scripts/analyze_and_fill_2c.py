import docx
import sys
import copy

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('bosung/lylich/2C_converted.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
for idx, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P{idx}: {p.text}")

print(f"\nTotal tables: {len(doc.tables)}")
for idx, t in enumerate(doc.tables):
    print(f"\n--- TABLE {idx}: rows={len(t.rows)}, cols={len(t.columns)} ---")
    for r_idx, r in enumerate(t.rows):
        cells_t = [c.text.strip().replace('\n', ' ') for c in r.cells]
        print(f"  R{r_idx}: {' | '.join(cells_t[:4])}")
