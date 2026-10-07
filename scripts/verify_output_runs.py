import sys
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'e:\OneDrive\HSCV\Antigravity\vbaibot\scripts\test_phieu_trinh_exact_output.docx')
t2 = doc.tables[2]

for pi in [1, 2, 9, 10, 20]:
    p = t2.rows[0].cells[0].paragraphs[pi]
    print(f'=== P{pi} ===')
    print('Full text:', p.text)
    for ri, r in enumerate(p.runs):
        if r.text:
            print(f'  run {ri}: text="{r.text}", bold={r.bold}')
