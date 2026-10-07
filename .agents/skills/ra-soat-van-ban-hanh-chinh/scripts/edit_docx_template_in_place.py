#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
edit_docx_template_in_place.py — Bộ công cụ chuẩn hóa chỉnh sửa file mẫu DOCX tại chỗ
NGUYÊN TẮC BẢO TOÀN CẤU TRÚC KHUNG 100% (ZERO STRUCTURE MUTATION ENGINE):
1. Giữ nguyên 100% table borders, cell dimensions, nested tables, margins, grid spans.
2. Giữ nguyên 100% font name, font size, bold, italic, line spacing, alignments của từng Run.
3. Thay thế nội dung ở cấp độ Run (Run-Level / XML-Level) — tuyệt đối không gán `paragraph.text = ...`
   làm xóa sạch `<w:r>` và mất trắng định dạng của tài liệu.
4. Hỗ trợ:
   - Exact text replacement (chuỗi trong 1 run hoặc kéo dài qua nhiều run)
   - Prefix-based replacement (thay giá trị sau một tiền tố/nhãn mà vẫn giữ nguyên format nhãn)
   - Khử trùng lặp (de-duplicate) cho các ô gộp (merged cells) trong bảng Word.
"""

import os
import sys
import json
import argparse
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')


def replace_text_in_paragraph_exact(paragraph, search_text, replace_text):
    """
    Thay thế search_text thành replace_text trên danh sách runs của paragraph
    mà TUYỆT ĐỐI KHÔNG làm mất font, size, bold, italic của các runs xung quanh.
    """
    if not search_text or search_text not in paragraph.text:
        return False

    full_text = "".join(r.text for r in paragraph.runs)
    if search_text not in full_text:
        return False

    # Tìm vị trí bắt đầu
    start_idx = full_text.find(search_text)
    end_idx = start_idx + len(search_text)

    # Tính toán tọa độ của từng run
    run_spans = []
    curr = 0
    for r in paragraph.runs:
        r_len = len(r.text)
        run_spans.append((curr, curr + r_len, r))
        curr += r_len

    # Tìm các run giao với khoảng [start_idx, end_idx)
    first_run_idx = None
    last_run_idx = None
    for i, (r_start, r_end, r) in enumerate(run_spans):
        if r_end > start_idx and r_start < end_idx:
            if first_run_idx is None:
                first_run_idx = i
            last_run_idx = i

    if first_run_idx is None:
        return False

    if first_run_idx == last_run_idx:
        # Chuỗi nằm trọn vẹn trong 1 run duy nhất
        r = paragraph.runs[first_run_idx]
        r_start, r_end, _ = run_spans[first_run_idx]
        offset_in_run = start_idx - r_start
        r.text = r.text[:offset_in_run] + replace_text + r.text[offset_in_run + len(search_text):]
    else:
        # Chuỗi kéo dài qua nhiều runs
        first_r_start, first_r_end, first_r = run_spans[first_run_idx]
        last_r_start, last_r_end, last_r = run_spans[last_run_idx]

        prefix = first_r.text[:(start_idx - first_r_start)]
        suffix = last_r.text[(end_idx - last_r_start):]

        first_r.text = prefix + replace_text
        last_r.text = suffix

        # Làm trống các run trung gian
        for mid_idx in range(first_run_idx + 1, last_run_idx):
            paragraph.runs[mid_idx].text = ""

    return True


def replace_after_prefix_in_paragraph(paragraph, prefix, new_value):
    """
    Tìm paragraph có chứa prefix (ví dụ: '1. Văn bản/cơ quan, đơn vị:'),
    giữ nguyên vẹn toàn bộ runs cấu thành prefix (kể cả bold/italic/font),
    và thay thế phần nội dung nằm sau prefix bằng new_value.
    """
    full_text = paragraph.text
    if prefix not in full_text:
        return False

    pos = full_text.find(prefix)
    after_prefix = full_text[pos + len(prefix):]
    stripped_old = after_prefix.lstrip()
    clean_new = str(new_value).strip()

    if stripped_old:
        return replace_text_in_paragraph_exact(paragraph, stripped_old, clean_new)
    else:
        # Nếu sau prefix chưa có text hoặc chỉ có khoảng trắng
        if paragraph.runs:
            # Thêm text vào run kế thừa sau cùng
            last_run = paragraph.runs[-1]
            if not last_run.text.endswith(" "):
                last_run.text += " " + clean_new
            else:
                last_run.text += clean_new
        else:
            paragraph.add_run(" " + clean_new)
        return True


def process_document_replacements(doc, exact_replacements=None, prefix_replacements=None):
    """
    Áp dụng bảng tra thay thế vào toàn bộ tài liệu:
    - exact_replacements: dict { search_text: replace_text }
    - prefix_replacements: dict { prefix_label: new_value }
    """
    if exact_replacements is None:
        exact_replacements = {}
    if prefix_replacements is None:
        prefix_replacements = {}

    count = 0

    def process_paragraph(p):
        nonlocal count
        # 1. Thay thế theo tiền tố (prefix) trước
        for prefix, new_val in prefix_replacements.items():
            if prefix in p.text:
                if replace_after_prefix_in_paragraph(p, prefix, str(new_val)):
                    count += 1

        # 2. Thay thế chính xác chuỗi (exact match)
        for s_text, r_text in exact_replacements.items():
            if s_text in p.text:
                # Lặp để thay thế hết các lần xuất hiện
                while s_text in p.text:
                    if replace_text_in_paragraph_exact(p, s_text, str(r_text)):
                        count += 1
                    else:
                        break

    # Duyệt paragraphs ngoài bảng
    for p in doc.paragraphs:
        process_paragraph(p)

    # Duyệt các Table (xử lý de-duplicate cho merged cells)
    processed_cells = set()

    def process_table(table):
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in processed_cells:
                    continue
                processed_cells.add(cell._tc)

                for p in cell.paragraphs:
                    process_paragraph(p)

                # Bảng lồng nhau
                for nested_table in cell.tables:
                    process_table(nested_table)

    for table in doc.tables:
        process_table(table)

    return count


def edit_docx_in_place(template_path, output_path, exact_replacements=None, prefix_replacements=None):
    """
    Thực hiện chỉnh sửa trực tiếp trên file mẫu DOCX và lưu ra output_path.
    Đảm bảo 100% nguyên bản về khung viền, bảng, font chữ.
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Không tìm thấy file mẫu: {template_path}")

    doc = docx.Document(template_path)
    count = process_document_replacements(doc, exact_replacements, prefix_replacements)
    doc.save(output_path)
    print(f"✅ Đã cập nhật {count} vị trí và xuất file thành công: {output_path}")
    return output_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Chỉnh sửa file DOCX mẫu tại chỗ bảo toàn 100% cấu trúc khung và font chữ")
    parser.add_argument("--template", required=True, help="Đường dẫn file DOCX mẫu gốc")
    parser.add_argument("--output", required=True, help="Đường dẫn lưu file DOCX kết quả")
    parser.add_argument("--exact", help="JSON string hoặc đường dẫn file JSON { search: replace }")
    parser.add_argument("--prefix", help="JSON string hoặc đường dẫn file JSON { prefix: new_value }")

    args = parser.parse_args()

    exact_data = {}
    if args.exact:
        if os.path.exists(args.exact):
            with open(args.exact, "r", encoding="utf-8") as f:
                exact_data = json.load(f)
        else:
            exact_data = json.loads(args.exact)

    prefix_data = {}
    if args.prefix:
        if os.path.exists(args.prefix):
            with open(args.prefix, "r", encoding="utf-8") as f:
                prefix_data = json.load(f)
        else:
            prefix_data = json.loads(args.prefix)

    edit_docx_in_place(args.template, args.output, exact_data, prefix_data)
