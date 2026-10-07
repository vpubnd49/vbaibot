#!/usr/bin/env python
# -*- coding: utf-8 -*-

import sys
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def replace_text_in_paragraph_exact(paragraph, search_text, replace_text):
    """
    Thay thế search_text thành replace_text trên danh sách runs của paragraph
    mà TUYỆT ĐỐI KHÔNG làm mất font, size, bold, italic của các runs xung quanh.
    """
    if not search_text:
        return False
        
    full_text = "".join(r.text for r in paragraph.runs)
    if search_text not in full_text:
        return False

    # Tìm vị trí bắt đầu
    start_idx = full_text.find(search_text)
    end_idx = start_idx + len(search_text)

    # Tính toán tọa độ của từng run
    run_spans = []
    curr = 0
    for r in paragraph.runs:
        r_len = len(r.text)
        run_spans.append((curr, curr + r_len, r))
        curr += r_len

    # Tìm các run giao với khoảng [start_idx, end_idx)
    first_run_idx = None
    last_run_idx = None
    for i, (r_start, r_end, r) in enumerate(run_spans):
        if r_end > start_idx and r_start < end_idx:
            if first_run_idx is None:
                first_run_idx = i
            last_run_idx = i

    if first_run_idx is None:
        return False

    if first_run_idx == last_run_idx:
        # Chuỗi nằm trọn vẹn trong 1 run
        r = paragraph.runs[first_run_idx]
        r_start, r_end, _ = run_spans[first_run_idx]
        offset_in_run = start_idx - r_start
        r.text = r.text[:offset_in_run] + replace_text + r.text[offset_in_run + len(search_text):]
    else:
        # Chuỗi kéo dài qua nhiều runs
        first_r_start, first_r_end, first_r = run_spans[first_run_idx]
        last_r_start, last_r_end, last_r = run_spans[last_run_idx]

        # Phần đầu của first_run trước search_text
        prefix = first_r.text[:(start_idx - first_r_start)]
        # Phần đuôi của last_run sau search_text
        suffix = last_r.text[(end_idx - last_r_start):]

        first_r.text = prefix + replace_text
        last_r.text = suffix

        # Các run ở giữa làm trống
        for mid_idx in range(first_run_idx + 1, last_run_idx):
            paragraph.runs[mid_idx].text = ""

    return True

doc = docx.Document(r'e:\OneDrive\HSCV\Antigravity\vbaibot\bosung\Phiếu trình.docx')
t2 = doc.tables[2]
p1 = t2.rows[0].cells[0].paragraphs[1]

print("BEFORE:")
print("p1 text:", p1.text)
print("p1 run 0:", p1.runs[0].text, "| bold:", p1.runs[0].bold)
if len(p1.runs) > 1:
    print("p1 run 1:", p1.runs[1].text, "| bold:", p1.runs[1].bold)

res = replace_text_in_paragraph_exact(
    p1, 
    "Tờ trình số 582/TTr-SNV ngày 26/11/2025 của Sở Nội vụ.",
    "Công văn số 4329/VPCP-CĐS ngày 14/9/2026 của Văn phòng Chính phủ."
)

print("\nREPLACE RESULT:", res)
print("AFTER:")
print("p1 text:", p1.text)
print("p1 run 0:", p1.runs[0].text, "| bold:", p1.runs[0].bold)
print("p1 run 1:", p1.runs[1].text, "| bold:", p1.runs[1].bold)
